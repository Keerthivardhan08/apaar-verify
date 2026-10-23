"""
Script to generate the complete, comprehensive APAAR Verify Demonstration Walkthrough
in both Microsoft Word (.docx) and Adobe PDF (.pdf) formats.
"""
import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


def set_cell_background(cell, fill_hex):
    """Sets background color of a docx table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def create_docx(filename):
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("APAAR Verify")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(16)
    run_sub = subtitle_p.add_run("Platform Demonstration & End-to-End Walkthrough Guide (v1.0)")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = RGBColor(0x43, 0x38, 0xCA)

    # Metadata banner table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    for row in meta_table.rows:
        for cell in row.cells:
            set_cell_background(cell, "F8FAFC")
            cell.width = Inches(3.4)

    meta_table.rows[0].cells[0].paragraphs[0].add_run("Product: APAAR Academic Integrity Platform").bold = True
    meta_table.rows[0].cells[1].paragraphs[0].add_run("Delivery: Production-Grade MVP Complete").bold = True
    meta_table.rows[1].cells[0].paragraphs[0].add_run("Target Audience: AICTE Reviewers, Institutions, Verifiers")
    meta_table.rows[1].cells[1].paragraphs[0].add_run("Architecture: FastAPI + React + RS256 PKI + Graph")

    doc.add_paragraph()

    # Disclaimer Callout Box
    disc_table = doc.add_table(rows=1, cols=1)
    disc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = disc_table.rows[0].cells[0]
    cell.width = Inches(6.8)
    set_cell_background(cell, "FEF3C7")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("⚠️ DEMONSTRATION & PROTOTYPE NOTICE:\n")
    r.bold = True
    r.font.color.rgb = RGBColor(0x92, 0x40, 0x0E)
    r2 = p.add_run(
        "APAAR Verify is an architectural and engineering prototype designed for the Smart India Hackathon. "
        "All institutional figures, student records, and verification responses shown in this demonstration are synthetic. "
        "Identity queries are simulated and do NOT constitute an official Ministry of Education or AICTE determination."
    )
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(0x78, 0x35, 0x0F)

    doc.add_paragraph()

    # Section 1: Credentials
    h1 = doc.add_heading("1. Demonstration Credentials & Role Matrix", level=1)
    h1.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_roles = doc.add_paragraph("The platform provides role-based authentication with strict multi-tenant isolation:")

    role_table = doc.add_table(rows=5, cols=4)
    role_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Role", "Email Username", "Password", "Primary Responsibilities"]
    col_widths = [Inches(1.5), Inches(1.8), Inches(1.1), Inches(2.4)]

    for idx, heading in enumerate(headers):
        cell = role_table.rows[0].cells[idx]
        cell.width = col_widths[idx]
        set_cell_background(cell, "1E3A8A")
        run = cell.paragraphs[0].add_run(heading)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    creds = [
        ("AICTE Admin", "admin@aicte.gov.in", "Demo@1234", "National oversight, intake ceilings, graph analytics, risk studio"),
        ("AICTE Reviewer", "reviewer@aicte.gov.in", "Demo@1234", "Investigate flagged anomalies, make defensible decisions, audit reviews"),
        ("IIT Demo Operator", "operator@iit-demo.edu", "Demo@1234", "Batch CSV uploads, student record management, draft corrections"),
        ("NIT Demo Operator", "operator@nit-demo.edu", "Demo@1234", "Independent institutional silo, separate student registry & batch processing"),
    ]

    for row_idx, data in enumerate(creds, start=1):
        for col_idx, text in enumerate(data):
            cell = role_table.rows[row_idx].cells[col_idx]
            cell.width = col_widths[col_idx]
            set_cell_background(cell, "FFFFFF" if row_idx % 2 == 1 else "F8FAFC")
            cell.paragraphs[0].add_run(text)

    doc.add_paragraph()

    # Section 2: End-to-End Scenarios
    h2 = doc.add_heading("2. Step-by-Step Interactive Demonstration Scenarios", level=1)
    h2.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    scenarios = [
        {
            "num": "Scenario 1",
            "title": "Institutional Batch CSV Ingestion & 9-Rule Deterministic Flagging",
            "actor": "Institution Operator (operator@iit-demo.edu)",
            "route": "http://localhost:5173/institution/upload",
            "api": "POST /api/v1/uploads",
            "steps": [
                "Log into the Institution Portal as operator@iit-demo.edu.",
                "Navigate to the 'Upload Student Data' workspace.",
                "Upload a synthetic admission batch CSV (e.g., synthetic_students_50.csv).",
                "Observe the 2-Stage Staging Pipeline: the system validates format, encoding, and schema before database commitment.",
                "The 9-rule integrity engine automatically flags seeded irregularities (e.g. APAAR_FORMAT_INVALID, REQUIRED_FIELD_MISSING, CATEGORY_INCONSISTENCY).",
                "Review the row-level error table, download the sanitized errors report (with formula injection shielding), and commit valid records.",
            ],
            "outcome": "Batch committed with full cryptographic checksum and row-level issue linkage in database.",
        },
        {
            "num": "Scenario 2",
            "title": "Multi-Stage Expanded Correction Reconciliation Workflow",
            "actor": "Operator -> Institution Admin -> AICTE Reviewer",
            "route": "http://localhost:5173/institution/corrections & /aicte/reviews",
            "api": "POST /api/v1/corrections-workflow",
            "steps": [
                "From the Student Detail view (/institution/students/:id), the Operator initiates an expanded correction request with proposed field changes and audit justification.",
                "The request enters stage 'pending_institution_admin'.",
                "Institution Admin reviews the institutional rationale, adds audit notes, and endorses it to 'pending_aicte_review'.",
                "The AICTE Reviewer accesses the National Review Queue (/aicte/reviews), compares side-by-side historical snapshots, and approves the change.",
                "The system creates an immutable StudentRecordVersion snapshot and automatically updates the active student record.",
            ],
            "outcome": "Complete multi-tiered regulatory audit trail with cryptographic versioning.",
        },
        {
            "num": "Scenario 3",
            "title": "Interactive Admission Integrity Graph & 30-Day Regulatory SLA",
            "actor": "AICTE Administrator (admin@aicte.gov.in)",
            "route": "http://localhost:5173/aicte/graph & /aicte/integrity",
            "api": "GET /api/v1/graph/topology & GET /api/v1/freshness/overview",
            "steps": [
                "Navigate to the Admission Integrity Graph (/aicte/graph).",
                "Inspect the interactive visual topology connecting Institutions, Programs, Cycles, and Student Records.",
                "Observe color-coded duplicate clusters linking students enrolled across multiple institutions.",
                "Navigate to the Freshness & SLA Dashboard to monitor institutional submission timestamps against the mandatory 30-day regulatory compliance window.",
            ],
            "outcome": "Instant visual detection of multi-enrollment syndicates and non-compliant stale submissions.",
        },
        {
            "num": "Scenario 4",
            "title": "Enterprise Asymmetric RS256 PKI Issuance & QR Verification",
            "actor": "Public Verifier / Employer / Student",
            "route": "http://localhost:5173/verify",
            "api": "GET /api/v1/.well-known/jwks.json & POST /api/v1/verification/verify-token",
            "steps": [
                "From a verified student profile, issue a Purpose-Bound Verifiable Credential (e.g. 'employment_screening' or 'admission_verification').",
                "The backend signs an asymmetric RSA-2048 token with SHA-256 (RS256) and embeds the Key ID (kid).",
                "Inspect the RFC 7517 dynamic JWKS endpoint at http://localhost:8000/api/v1/.well-known/jwks.json.",
                "Navigate to the public verification portal (/verify).",
                "Scan the generated QR code or paste the cryptographic token to execute offline verification against the public key.",
                "Observe the strict verification result: VALID status, issuer confirmation, and zero PII leakage beyond permitted purpose.",
            ],
            "outcome": "Zero-trust verification without exposing sensitive student identity numbers.",
        },
        {
            "num": "Scenario 5",
            "title": "Direct Campus ERP Ingestion & Blinded Zero-Knowledge Exchange",
            "actor": "Automated Campus SIS & Verification Agent",
            "route": "REST API Integration & /institution/overview",
            "api": "POST /api/v1/institution-api/v1/sis/push-records & POST /api/v1/privacy-exchange/blinded-proof",
            "steps": [
                "Generate a cryptographically hashed Institutional API Key with permission scoping.",
                "Simulate Campus ERP batch ingestion using HTTP header 'X-Institution-API-Key: inst_live_...'.",
                "Trigger a Zero-Knowledge Blinded Proof request for predicate 'CURRENTLY_ENROLLED_REGULAR_STUDENT'.",
                "Verify the blinded proof token without decrypting or storing any underlying personal data.",
            ],
            "outcome": "Seamless machine-to-machine sync and cryptographically blinded mathematical assertions.",
        },
        {
            "num": "Scenario 6",
            "title": "OCR-Assisted Document & Marksheet Reconciliation",
            "actor": "Institution Operator / Reviewer",
            "route": "http://localhost:5173/institution/students/:id (OCR Tab)",
            "api": "POST /api/v1/ocr/reconcile-document",
            "steps": [
                "Open any student profile and select the new 'OCR Document Reconciliation' tab.",
                "Click 'Simulate Valid Marksheet' to pre-populate scanned certificate fields (Candidate Name, Roll No, DOB, Marks, Board).",
                "Click 'Execute OCR Fuzzy Reconciliation' to observe instant fuzzy string matching and 98%+ confidence score ('VERIFIED').",
                "Click 'Simulate Conflicting Document' with mismatched candidate names or altered roll numbers.",
                "Execute reconciliation to observe automatic 'MISMATCH_FLAGGED' status and automated creation of a DOCUMENT_DATA_MISMATCH compliance flag.",
            ],
            "outcome": "Automated fraud prevention against tampered or conflicting paper credentials.",
        },
        {
            "num": "Scenario 7",
            "title": "National Integrity Risk Studio & 30-Day Anomaly Forecasting",
            "actor": "AICTE Administrator (admin@aicte.gov.in)",
            "route": "http://localhost:5173/aicte/analytics",
            "api": "GET /api/v1/analytics/risk-scores & GET /api/v1/analytics/trends",
            "steps": [
                "Navigate to National Analytics (/aicte/analytics).",
                "Switch to the 'Risk Scoring Studio' tab to inspect institutional composite risk scores (0-100) combining anomaly density, severity weighting, duplicate clustering, and capacity breaches.",
                "Review Risk Tiers (Critical, High, Moderate, Low Risk) and automated AICTE regulatory audit recommendations.",
                "Switch to 'Model Performance' to view precision rates (96%+ average) and false-positive rates across all 9 rule engines.",
            ],
            "outcome": "Macro-level statistical forecasting and automated risk prioritization for regulatory inspections.",
        },
    ]

    for sc in scenarios:
        doc.add_heading(f"{sc['num']}: {sc['title']}", level=2)
        p_meta = doc.add_paragraph()
        p_meta.add_run(f"Primary Actor: ").bold = True
        p_meta.add_run(f"{sc['actor']}  |  ")
        p_meta.add_run(f"UI Route: ").bold = True
        p_meta.add_run(f"{sc['route']}\n")
        p_meta.add_run(f"API Endpoint: ").bold = True
        p_meta.add_run(f"{sc['api']}")

        doc.add_paragraph("Action Steps:").bold = True
        for idx, step in enumerate(sc['steps'], start=1):
            p_step = doc.add_paragraph(f"{idx}. {step}")
            p_step.paragraph_format.left_indent = Inches(0.2)
            p_step.paragraph_format.space_after = Pt(2)

        p_out = doc.add_paragraph()
        p_out.paragraph_format.space_before = Pt(4)
        r_out_lbl = p_out.add_run("Expected Outcome: ")
        r_out_lbl.bold = True
        r_out_lbl.font.color.rgb = RGBColor(0x16, 0x65, 0x34)
        p_out.add_run(sc['outcome'])
        doc.add_paragraph()

    # Section 3: Shareable Demo Link & Networking Instructions
    h3 = doc.add_heading("3. Shareable Link & Live Demo Access Instructions", level=1)
    h3.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    doc.add_paragraph(
        "To present or share the live demonstration with team members, evaluators, or public users:"
    )

    share_table = doc.add_table(rows=5, cols=3)
    share_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_headers = ["Service / Portal", "Local Access URL", "External / Shareable Access"]
    s_widths = [Inches(2.2), Inches(2.2), Inches(2.4)]

    for idx, heading in enumerate(s_headers):
        cell = share_table.rows[0].cells[idx]
        cell.width = s_widths[idx]
        set_cell_background(cell, "1E3A8A")
        r = cell.paragraphs[0].add_run(heading)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    share_rows = [
        ("Main Web Application", "http://localhost:5173", "https://7f4fee2bc4e56f.lhr.life (Direct Public Tunnel)"),
        ("Public Credential Verifier", "http://localhost:5173/verify", "https://7f4fee2bc4e56f.lhr.life/verify"),
        ("FastAPI Interactive Docs", "http://localhost:8000/api/v1/docs", "Live OpenAPI specification & testing"),
        ("Dynamic JWKS Public Keys", "http://localhost:8000/api/v1/.well-known/jwks.json", "https://7f4fee2bc4e56f.lhr.life/api/v1/.well-known/jwks.json"),
    ]

    for row_idx, data in enumerate(share_rows, start=1):
        for col_idx, text in enumerate(data):
            cell = share_table.rows[row_idx].cells[col_idx]
            cell.width = s_widths[col_idx]
            set_cell_background(cell, "FFFFFF" if row_idx % 2 == 1 else "F8FAFC")
            cell.paragraphs[0].add_run(text)

    doc.add_paragraph()

    doc.add_heading("Starting the Services for Local Network Sharing:", level=3)
    doc.add_paragraph(
        "1. Start the FastAPI backend with 0.0.0.0 binding:\n"
        "   cd backend && python -m uvicorn app.main:app --host 0.0.0.0 --port 8000 --reload\n\n"
        "2. Start the Vite React frontend with host exposure:\n"
        "   cd frontend && npm run dev -- --host 0.0.0.0 --port 5173\n\n"
        "3. Optional Cloud Tunnel for external internet sharing:\n"
        "   npx localtunnel --port 5173\n"
        "   (or run: ngrok http 5173)"
    )

    doc.save(filename)
    print(f"Successfully generated DOCX at: {filename}")


def create_pdf(filename):
    pdf = SimpleDocTemplate(
        filename,
        pagesize=letter,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45,
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=22,
        leading=26,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=4,
    )

    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#4338CA'),
        spaceAfter=14,
    )

    h1_style = ParagraphStyle(
        'DocH1',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#1E3A8A'),
        spaceBefore=12,
        spaceAfter=6,
    )

    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading3'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#1E3A8A'),
        spaceBefore=8,
        spaceAfter=4,
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#334155'),
        spaceAfter=4,
    )

    step_style = ParagraphStyle(
        'DocStep',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#1E293B'),
        leftIndent=14,
        spaceAfter=2,
    )

    disclaimer_style = ParagraphStyle(
        'DocDisclaimer',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor('#78350F'),
    )

    elements = []

    # Title & Subtitle
    elements.append(Paragraph("APAAR Verify", title_style))
    elements.append(Paragraph("Platform Demonstration & End-to-End Walkthrough Guide (v1.0)", subtitle_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#1E3A8A'), spaceAfter=10))

    # Disclaimer Callout
    disc_text = (
        "<b>⚠️ DEMONSTRATION & PROTOTYPE NOTICE:</b><br/>"
        "APAAR Verify is an architectural and engineering prototype designed for the Smart India Hackathon. "
        "All institutional figures, student records, and verification responses shown in this demonstration are synthetic. "
        "Identity queries are simulated and do NOT constitute an official Ministry of Education or AICTE determination."
    )
    disc_table = Table([[Paragraph(disc_text, disclaimer_style)]], colWidths=[522])
    disc_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#FEF3C7')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#F59E0B')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
        ('RIGHTPADDING', (0, 0), (-1, -1), 10),
    ]))
    elements.append(disc_table)
    elements.append(Spacer(1, 10))

    # Section 1: Credentials
    elements.append(Paragraph("1. Demonstration Credentials & Role Matrix", h1_style))
    elements.append(Paragraph("The platform enforces strict role-based access control (RBAC) with institutional data isolation:", body_style))

    cred_data = [
        [Paragraph("<b>Role</b>", body_style), Paragraph("<b>Email Username</b>", body_style), Paragraph("<b>Password</b>", body_style), Paragraph("<b>Responsibilities</b>", body_style)],
        [Paragraph("AICTE Admin", body_style), Paragraph("admin@aicte.gov.in", body_style), Paragraph("Demo@1234", body_style), Paragraph("National monitoring, capacity ceilings, graph analytics", body_style)],
        [Paragraph("AICTE Reviewer", body_style), Paragraph("reviewer@aicte.gov.in", body_style), Paragraph("Demo@1234", body_style), Paragraph("Investigate anomalies, audit decisions", body_style)],
        [Paragraph("IIT Demo Operator", body_style), Paragraph("operator@iit-demo.edu", body_style), Paragraph("Demo@1234", body_style), Paragraph("Batch CSV uploads, draft corrections", body_style)],
        [Paragraph("NIT Demo Operator", body_style), Paragraph("operator@nit-demo.edu", body_style), Paragraph("Demo@1234", body_style), Paragraph("Independent institutional silo & uploads", body_style)],
    ]
    cred_table = Table(cred_data, colWidths=[105, 140, 80, 197])
    cred_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(cred_table)
    elements.append(Spacer(1, 12))

    # Section 2: Walkthrough Scenarios
    elements.append(Paragraph("2. Step-by-Step Interactive Demonstration Scenarios", h1_style))

    scenarios = [
        ("Scenario 1: Institutional Batch Ingestion & 9-Rule Flagging",
         "Operator (operator@iit-demo.edu) | UI: /institution/upload",
         [
             "Log into Institution Portal as operator@iit-demo.edu.",
             "Navigate to 'Upload Student Data' workspace.",
             "Upload synthetic admission batch CSV; observe the 2-Stage Staging Pipeline.",
             "The 9-rule integrity engine automatically flags seeded errors (APAAR_FORMAT_INVALID, REQUIRED_FIELD_MISSING, CATEGORY_INCONSISTENCY).",
             "Review row-level error reports, download sanitized CSV with formula injection shielding, and commit valid records.",
         ],
         "Batch committed with cryptographic checksum and row-level issue linkage."),

        ("Scenario 2: Multi-Stage Correction Reconciliation Workflow",
         "Operator -> Inst Admin -> AICTE Reviewer | UI: /institution/corrections & /aicte/reviews",
         [
             "From Student Detail (/institution/students/:id), Operator submits correction with proposed value & justification.",
             "Inst Admin reviews institutional rationale and endorses request to AICTE review.",
             "AICTE Reviewer accesses Review Queue, compares side-by-side snapshots, and confirms or adjusts decision.",
             "System generates immutable StudentRecordVersion snapshot and applies update to active record.",
         ],
         "Defensible regulatory audit trail with cryptographic versioning."),

        ("Scenario 3: Admission Integrity Graph & 30-Day Freshness SLA",
         "AICTE Admin (admin@aicte.gov.in) | UI: /aicte/graph & /aicte/integrity",
         [
             "Open Admission Integrity Graph (/aicte/graph) to visualize relationships across Institutions, Programs, and Students.",
             "Inspect duplicate syndicate clusters spanning multiple institutions in real time.",
             "Navigate to Source Freshness Dashboard to verify submission compliance against the 30-day regulatory SLA.",
         ],
         "Instant visual detection of multi-enrollment rings and SLA delinquency."),

        ("Scenario 4: Asymmetric RS256 PKI Issuance & QR Verification",
         "Public Verifier / Student | UI: /verify | JWKS: /.well-known/jwks.json",
         [
             "Issue a Purpose-Bound Verifiable Credential from student profile (e.g. employment_screening).",
             "Backend signs RSA-2048 token with RS256 algorithm and Key ID (kid).",
             "Inspect RFC 7517 dynamic JWKS endpoint at /api/v1/.well-known/jwks.json.",
             "Scan QR code or paste token into public verifier portal (/verify) for zero-trust offline mathematical verification.",
         ],
         "Verifiable credential validation with zero raw PII disclosure."),

        ("Scenario 5: Direct Campus ERP Ingestion & Blinded ZK Proofs",
         "Campus SIS Gateway & Verifier | REST API Integration",
         [
             "Generate hashed Institutional API Key (X-Institution-API-Key: inst_live_...).",
             "Push machine-to-machine records directly to /api/v1/institution-api/v1/sis/push-records.",
             "Request Blinded Zero-Knowledge Proof for predicate CURRENTLY_ENROLLED_REGULAR_STUDENT.",
             "Verify blinded assertion token without revealing student identity or marks.",
         ],
         "Automated enterprise sync and zero-knowledge claim assertions."),

        ("Scenario 6: OCR-Assisted Document & Marksheet Reconciliation",
         "Operator / Reviewer | UI: /institution/students/:id (OCR Tab)",
         [
             "Open student profile and click the new 'OCR Document Reconciliation' tab.",
             "Click 'Simulate Valid Marksheet' to auto-populate extracted certificate attributes.",
             "Execute reconciliation to observe fuzzy string similarity and 98%+ confidence score (VERIFIED).",
             "Click 'Simulate Conflicting Document' to test tampered names or conflicting roll numbers.",
             "Observe automatic MISMATCH_FLAGGED status and automatic creation of a DOCUMENT_DATA_MISMATCH flag.",
         ],
         "Automated detection of fraudulent and tampered paper documents."),

        ("Scenario 7: National Integrity Risk Studio & Anomaly Forecasting",
         "AICTE Admin (admin@aicte.gov.in) | UI: /aicte/analytics",
         [
             "Navigate to National Analytics (/aicte/analytics) and click 'Risk Scoring Studio'.",
             "Inspect institutional composite risk rankings (0-100) based on anomaly density, severity weighting, duplicate clusters, and capacity excess.",
             "Review Risk Tiers (Critical, High, Moderate, Low Risk) and automated regulatory audit directives.",
             "Switch to 'Model Performance' tab to review precision benchmarks (96%+ average) across all 9 validation rules.",
         ],
         "Macro-level predictive modeling and risk-prioritized regulatory review."),
    ]

    for title, meta, steps, outcome in scenarios:
        elements.append(Paragraph(title, h2_style))
        elements.append(Paragraph(f"<b>Context:</b> {meta}", body_style))
        for idx, s in enumerate(steps, 1):
            elements.append(Paragraph(f"{idx}. {s}", step_style))
        elements.append(Paragraph(f"<b>Expected Outcome:</b> <font color='#166534'>{outcome}</font>", body_style))
        elements.append(Spacer(1, 4))

    # Section 3: Shareable Demo Link
    elements.append(Spacer(1, 8))
    elements.append(Paragraph("3. Shareable Link & Live Demo Access Instructions", h1_style))
    elements.append(Paragraph("To access or share the demonstration with team members and evaluators:", body_style))

    share_data = [
        [Paragraph("<b>Service / Portal</b>", body_style), Paragraph("<b>Local Access URL</b>", body_style), Paragraph("<b>External / Shareable Access</b>", body_style)],
        [Paragraph("Main Web App", body_style), Paragraph("http://localhost:5173", body_style), Paragraph("https://7f4fee2bc4e56f.lhr.life", body_style)],
        [Paragraph("Public Verifier", body_style), Paragraph("http://localhost:5173/verify", body_style), Paragraph("https://7f4fee2bc4e56f.lhr.life/verify", body_style)],
        [Paragraph("FastAPI Swagger Docs", body_style), Paragraph("http://localhost:8000/api/v1/docs", body_style), Paragraph("Interactive API testing", body_style)],
        [Paragraph("Dynamic JWKS Endpoint", body_style), Paragraph("http://localhost:8000/api/v1/.well-known/jwks.json", body_style), Paragraph("https://7f4fee2bc4e56f.lhr.life/api/v1/.well-known/jwks.json", body_style)],
    ]
    share_table = Table(share_data, colWidths=[130, 192, 200])
    share_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(share_table)

    pdf.build(elements)
    print(f"Successfully generated PDF at: {filename}")


if __name__ == "__main__":
    docx_file = "APAAR_Verify_Demo_Walkthrough.docx"
    pdf_file = "APAAR_Verify_Demo_Walkthrough.pdf"

    create_docx(docx_file)
    create_pdf(pdf_file)

    # Also copy to artifact directory
    artifact_dir = r"C:\Users\Keerthi\.gemini\antigravity\brain\44e0a93b-ed11-4188-8c89-e5dfe2edb7ff"
    if os.path.exists(artifact_dir):
        shutil.copy(docx_file, os.path.join(artifact_dir, docx_file))
        shutil.copy(pdf_file, os.path.join(artifact_dir, pdf_file))
        print(f"Copied both documents to artifact directory: {artifact_dir}")
